"""Anything you drop in the meeting folder: an agenda, a deck, a one-pager.

Text only, and deliberately so. A meeting's minutes are made from what was said;
a document is context that stops the model mishearing a product name, not a second
source of decisions.
"""
from __future__ import annotations
import pathlib

MAX_CHARS = 60_000
OURS = {".md", ".pdf", ".tex"}


def gather(folder: pathlib.Path) -> tuple[list[tuple[str, str]], list[str]]:
    out, skipped = [], []
    for p in sorted(folder.iterdir()):
        if not p.is_file() or p.name.startswith((".", "~$")):
            continue
        if p.suffix.lower() in OURS and p.stem.startswith(folder.name.split(" — ")[-1]):
            continue                                  # our own output
        text = extract(p).strip()
        if text:
            out.append((p.name, text[:MAX_CHARS]))
        else:
            skipped.append(p.name)
    return out, skipped


def extract(path: pathlib.Path) -> str:
    s = path.suffix.lower()
    try:
        if s == ".pdf":
            import pypdf
            return "\n".join((pg.extract_text() or "")
                             for pg in pypdf.PdfReader(str(path)).pages)
        if s == ".docx":
            import docx
            return "\n".join(par.text for par in docx.Document(str(path)).paragraphs)
        if s == ".pptx":
            from pptx import Presentation
            bits = []
            for i, slide in enumerate(Presentation(str(path)).slides, 1):
                texts = [sh.text for sh in slide.shapes if getattr(sh, "has_text_frame", False)]
                if texts:
                    bits.append(f"--- slide {i} ---\n" + "\n".join(texts))
            return "\n".join(bits)
        if s in {".txt", ".md", ".csv", ".tex"}:
            return path.read_text(encoding="utf-8", errors="replace")
    except Exception:
        return ""
    return ""
